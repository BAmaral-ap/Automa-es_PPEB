from contextlib import redirect_stdout, redirect_stderr
from datetime import datetime, time
import io
from pathlib import Path
import sys
import tempfile
import unittest
from zipfile import ZipFile

from lxml import etree
from openpyxl import load_workbook

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
import run
from convite import CAIXA_CORPO, conferir_pdf_convite
from layout import CAIXAS, NS, conferir_pdf


def textos_modernos(tree):
    return [" ".join(p.xpath("./w:r/w:t/text()", namespaces=NS))
            for p in tree.xpath("//wp:anchor//w:txbxContent/w:p", namespaces=NS)]


def geometria(tree):
    return {a.find("wp:docPr", NS).get("name"):
            (a.find("wp:positionH", NS).get("relativeFrom"),
             a.find("wp:positionV", NS).get("relativeFrom"),
             a.find("wp:positionH/wp:posOffset", NS).text,
             a.find("wp:positionV/wp:posOffset", NS).text,
             dict(a.find("wp:extent", NS).attrib))
            for a in tree.xpath('//wp:anchor[starts-with(wp:docPr/@name,"PPEB_")]', namespaces=NS)}


class AutomacaoTest(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.addCleanup(self.temp.cleanup)
        self.root = Path(self.temp.name)
        self.planilha = self.root / "dados.xlsx"
        original = load_workbook(run.BASE / "dados_qualificacao_doutorado.xlsx")
        original.save(self.planilha)
        original.close()

    def alterar(self, coluna, valor):
        wb = load_workbook(self.planilha)
        ws = wb["Dados"]
        headers = [c.value for c in ws[1]]
        ws.cell(2, headers.index(coluna) + 1, valor)
        wb.save(self.planilha)
        wb.close()

    def executar(self, *extras):
        with redirect_stdout(io.StringIO()), redirect_stderr(io.StringIO()):
            return run.main(["--planilha", str(self.planilha), "--saida", str(self.root / "saida"),
                             "--formato", "docx", *extras])

    def test_geracao_preserva_imagens_e_estrutura(self):
        self.assertEqual(self.executar(), 0)
        docx = next((self.root / "saida").glob("*/Cartaz*.docx"))
        ns = {"w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main"}
        with ZipFile(docx) as rendered, ZipFile(run.BASE / "modelo_qualificacao_doutorado.docx") as template:
            self.assertIsNone(rendered.testzip())
            xml = rendered.read("word/document.xml")
            tree = etree.fromstring(xml)
            textos = textos_modernos(tree)
            self.assertIn("ADRIAN SOUZA DOS SANTOS", textos)
            self.assertIn("31/03/2026", textos)
            self.assertIn("14h00", textos)
            self.assertNotIn(b"{{", xml)
            self.assertEqual(sum(" - Membro Externo" in t for t in textos), 3)
            self.assertEqual(sum(" - Suplente" in t for t in textos), 2)
            media = [n for n in template.namelist() if n.startswith("word/media/")]
            self.assertTrue(media)
            for nome in media:
                self.assertEqual(template.read(nome), rendered.read(nome))
            before = etree.fromstring(template.read("word/document.xml"))
            self.assertEqual(geometria(before), geometria(tree))
            self.assertEqual(len(geometria(tree)), 7)
            self.assertFalse(tree.xpath("//w:tbl", namespaces=ns))

    def test_caracteres_especiais_nao_quebram_xml(self):
        titulo = "Educação & sociedade: <escola> {{ texto literal }}"
        self.alterar("Título", titulo)
        self.assertEqual(self.executar(), 0)
        with ZipFile(next((self.root / "saida").glob("*/Cartaz*.docx"))) as z:
            tree = etree.fromstring(z.read("word/document.xml"))
            self.assertIn(titulo, textos_modernos(tree))

    def test_campos_opcionais_nao_deixam_rotulos(self):
        wb = load_workbook(self.planilha)
        ws = wb["Dados"]
        for cell in ws[1]:
            if cell.value in ("Membro Externo 3", "Universidade Externo 3", "Suplente 2", "Universidade Suplente 2"):
                ws.cell(2, cell.column).value = None
        wb.save(self.planilha)
        wb.close()
        self.assertEqual(self.executar(), 0)
        with ZipFile(next((self.root / "saida").glob("*/Cartaz*.docx"))) as z:
            tree = etree.fromstring(z.read("word/document.xml"))
            texts = textos_modernos(tree)
            self.assertEqual(sum(" - Membro Externo" in t for t in texts), 2)
            self.assertEqual(sum(" - Suplente" in t for t in texts), 1)
            ns = {"w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main"}
            self.assertNotIn("• ", texts)
            self.assertFalse(tree.xpath("//w:pPr/w:numPr", namespaces=ns))
        with ZipFile(next((self.root / "saida").glob("*/Convite*.docx"))) as z:
            tree = etree.fromstring(z.read("word/document.xml"))
            corpo = tree.xpath('//wp:anchor[wp:docPr/@name="PPEB_convite"]//w:txbxContent', namespaces=NS)[0]
            textos = corpo.xpath('.//w:t/text()', namespaces=NS)
            self.assertEqual(sum(" - Membro Externo" in t for t in textos), 2)
            self.assertEqual(sum(" - Suplente" in t for t in textos), 1)
            self.assertTrue(all("".join(p.xpath('.//w:t/text()', namespaces=NS)).strip()
                                for p in corpo.findall('w:p', NS)))

    def test_convite_preenche_planilha_e_preserva_arte_e_assinatura(self):
        titulo = "Educação & sociedade: <escola> {{ texto literal }}"
        self.alterar("Título", titulo)
        self.assertEqual(self.executar(), 0)
        arquivos = list((self.root / "saida").glob("*/*"))
        self.assertEqual(sorted(p.suffix for p in arquivos), ['.docx', '.docx'])
        convite = next(p for p in arquivos if p.name.startswith('Convite'))
        with ZipFile(convite) as z, ZipFile(run.BASE / 'modelo_convite_doutorado.docx') as modelo:
            self.assertIsNone(z.testzip())
            tree = etree.fromstring(z.read('word/document.xml'))
            corpos = tree.xpath('//wp:anchor[wp:docPr/@name="PPEB_convite"]//w:txbxContent | '
                                '//v:shape[@id="PPEB_convite"]//w:txbxContent', namespaces=NS)
            self.assertEqual(len(corpos), 2)
            textos = ["".join(c.xpath('.//w:t/text()', namespaces=NS)) for c in corpos]
            self.assertEqual(textos[0], textos[1])
            for valor in (titulo, 'ADRIAN SOUZA DOS SANTOS', '31/03/2026', '14h00', 'Sala NEB-02 / NEB-UFPA'):
                self.assertIn(valor, textos[0])
            self.assertNotIn('VIVIAN DE LIMA CABRAL', textos[0])
            self.assertNotIn('20/02/2026', textos[0])
            self.assertNotIn('{%', textos[0])
            for nome in modelo.namelist():
                if nome.startswith('word/media/'):
                    self.assertEqual(z.read(nome), modelo.read(nome))
            original = etree.fromstring(modelo.read('word/document.xml'))
            assinaturas = []
            for t in (tree, original):
                partes = [c for c in t.xpath('//w:txbxContent', namespaces=NS)
                          if 'Atenciosamente' in ''.join(c.xpath('.//w:t/text()', namespaces=NS))]
                self.assertEqual(len(partes), 2)
                # Prefixos e declarações de namespace podem mudar sem alterar o XML.
                assinaturas.append([[(e.tag, dict(e.attrib), e.text) for e in c.iter()]
                                    for c in partes])
            self.assertEqual(assinaturas[0], assinaturas[1])

    def test_textos_longos_reduzem_fontes_sem_mover_caixas(self):
        self.alterar("Título", "Educação, currículo e políticas de gestão democrática nas escolas públicas. " * 5)
        self.alterar("Aluno", "ANA MARIA DE SOUZA DOS SANTOS DE ALMEIDA E SILVA")
        self.alterar("Orientador", "Profª. Drª. Maria Aparecida de Albuquerque de Souza e Silva")
        self.alterar("Sala", "Auditório do Núcleo de Estudos em Educação Básica, Campus Belém - UFPA")
        self.assertEqual(self.executar(), 0)
        with ZipFile(next((self.root / "saida").glob("*/Cartaz*.docx"))) as z, \
                ZipFile(run.BASE / "modelo_qualificacao_doutorado.docx") as original:
            tree = etree.fromstring(z.read("word/document.xml"))
            self.assertEqual(geometria(tree), geometria(etree.fromstring(original.read("word/document.xml"))))
            for chave in ("titulo", "aluno", "orientador", "sala"):
                sizes = tree.xpath(f'//wp:anchor[wp:docPr/@name="PPEB_{chave}"]//w:sz/@w:val', namespaces=NS)
                self.assertTrue(sizes)
                self.assertLess(max(map(int, sizes)), CAIXAS[chave].fonte * 2)
            for chave in CAIXAS:
                modern = tree.xpath(f'//wp:anchor[wp:docPr/@name="PPEB_{chave}"]//w:t/text()', namespaces=NS)
                fallback = tree.xpath(f'//v:shape[@id="PPEB_{chave}"]//w:t/text()', namespaces=NS)
                self.assertEqual(modern, fallback)

    def test_texto_impossivel_de_encaixar_falha_sem_cortar(self):
        self.alterar("Título", "Texto enorme " * 1000)
        self.assertEqual(self.executar(), 1)
        self.assertFalse(list((self.root / "saida").glob("*/*.docx")))

    def test_nao_sobrescreve_sem_opcao(self):
        self.assertEqual(self.executar(), 0)
        path = next((self.root / "saida").glob("*/Cartaz*.docx"))
        before = path.read_bytes()
        self.alterar("Título", "Título novo")
        self.assertEqual(self.executar(), 1)
        self.assertEqual(path.read_bytes(), before)
        self.assertEqual(self.executar("--sobrescrever"), 0)
        self.assertNotEqual(path.read_bytes(), before)

    def test_entrada_invalida_nao_produz_arquivos(self):
        self.alterar("Data", "31/02/2026")
        self.assertEqual(self.executar(), 1)
        self.assertFalse((self.root / "saida").exists())

    def test_formula_rejeitada(self):
        self.alterar("Título", '=CONCAT("a","b")')
        self.assertEqual(self.executar(), 1)

    def test_aluno_duplicado_rejeitado(self):
        wb = load_workbook(self.planilha)
        ws = wb["Dados"]
        ws.append([c.value for c in ws[2]])
        wb.save(self.planilha)
        wb.close()
        self.assertEqual(self.executar(), 1)
        self.assertFalse((self.root / "saida").exists())

    def test_formatacao_de_datas_horas_excel(self):
        self.assertEqual(run.formatar_data(datetime(2026, 3, 31)), "31/03/2026")
        self.assertEqual(run.formatar_hora(time(14)), "14h00")
        self.assertEqual(run.formatar_hora(14 / 24), "14h00")
        self.assertEqual(run.formatar_hora("9:05"), "09h05")
        with self.assertRaises(ValueError):
            run.formatar_hora("24h00")

    def test_nome_pasta_compativel_com_windows(self):
        self.assertEqual(run.nome_pasta("../Ana/João"), "_Ana_João")
        self.assertEqual(run.nome_pasta("CON"), "Aluno - CON")

    def test_pdf_multipagina_nao_vira_imagem_incompleta(self):
        import pymupdf
        path = self.root / "duas_paginas.pdf"
        with pymupdf.open() as pdf:
            pdf.new_page()
            pdf.new_page()
            pdf.save(path)
        with self.assertRaisesRegex(ValueError, "2 páginas"):
            run.gerar_png(path, 72)
        self.assertFalse(path.with_suffix(".png").exists())

    def test_falha_de_conversor_retorna_erro_e_nao_publica_docx(self):
        from unittest.mock import patch
        with patch.object(run, "escolher_conversor", return_value="libreoffice"), \
             patch.object(run, "gerar_pdf", side_effect=RuntimeError("Falha simulada")), \
             redirect_stdout(io.StringIO()), redirect_stderr(io.StringIO()):
            code = run.main(["--planilha", str(self.planilha), "--saida", str(self.root / "saida")])
        self.assertEqual(code, 1)
        self.assertFalse(list((self.root / "saida").glob("*/*.docx")))

    def test_falha_no_convite_nao_entrega_cartaz_parcial(self):
        from unittest.mock import patch
        with patch.object(run, 'gerar_convite', side_effect=RuntimeError('Convite inválido')):
            self.assertEqual(self.executar(), 1)
        self.assertFalse(list((self.root / 'saida').glob('*/*')))

    def test_modelo_convite_ausente_nao_produz_arquivos(self):
        self.assertEqual(self.executar('--modelo-convite', str(self.root / 'ausente.docx')), 1)
        self.assertFalse((self.root / 'saida').exists())

    @unittest.skipUnless(run.localizar_soffice(), "LibreOffice não disponível")
    def test_pdf_real_preserva_caixas_e_texto_curto_e_longo(self):
        import pymupdf
        variantes = [
            {"Título": "Educação e currículo"},
            {},  # Dados do documento enviado.
            {"Título": "Educação, currículo e políticas de gestão democrática nas escolas públicas. " * 5,
             "Aluno": "ANA MARIA DE SOUZA DOS SANTOS DE ALMEIDA E SILVA",
             "Orientador": "Profª. Drª. Maria Aparecida de Albuquerque de Souza e Silva",
             "Sala": "Auditório do Núcleo de Estudos em Educação Básica, Campus Belém - UFPA",
             "Membro Externo 1": "Profª. Drª. Maria Aparecida de Albuquerque de Souza e Silva de Almeida",
             "Membro Externo 2": "Prof. Dr. José de Albuquerque de Souza dos Santos de Almeida",
             "Universidade Externo 1": "Universidade Federal do Pará",
             "Universidade Externo 2": "Universidade Federal do Paraná"},
        ]
        for i, alteracoes in enumerate(variantes):
            original = load_workbook(run.BASE / "dados_qualificacao_doutorado.xlsx")
            original.save(self.planilha)
            original.close()
            for chave, valor in alteracoes.items():
                self.alterar(chave, valor)
            ctx = run.ler_planilha(self.planilha)[0][2]
            saida = self.root / f"renderizado-{i}"
            with redirect_stdout(io.StringIO()), redirect_stderr(io.StringIO()):
                code = run.main(["--planilha", str(self.planilha), "--saida", str(saida), "--dpi", "72"])
            self.assertEqual(code, 0, alteracoes)
            pdf = next(saida.glob("*/Cartaz*.pdf"))
            self.assertEqual(conferir_pdf(pdf, ctx), [])
            with pymupdf.open(pdf) as doc:
                self.assertEqual(len(doc), 1)
                # Verifica os retângulos efetivamente desenhados, além das âncoras do DOCX.
                rects = [d["rect"] for d in doc[0].get_drawings()]
                for caixa in CAIXAS.values():
                    target = pymupdf.Rect(caixa.x, caixa.y, caixa.x + caixa.largura, caixa.y + caixa.altura)
                    self.assertTrue(any(max(abs(a-b) for a,b in zip(rect, target)) < 1 for rect in rects), target)
            convite = next(saida.glob('*/Convite*.pdf'))
            with ZipFile(convite.with_suffix('.docx')) as z:
                tree = etree.fromstring(z.read('word/document.xml'))
                corpo = tree.xpath('//wp:anchor[wp:docPr/@name="PPEB_convite"]//w:txbxContent', namespaces=NS)[0]
                esperado = ' '.join(''.join(p.xpath('.//w:t/text()', namespaces=NS))
                                    for p in corpo.findall('w:p', NS))
            self.assertTrue(conferir_pdf_convite(convite, esperado), alteracoes)
            with pymupdf.open(convite) as doc:
                self.assertEqual(len(doc), 1)
                self.assertIn('Atenciosamente', doc[0].get_text())
                self.assertTrue(doc[0].get_images())
                # Corpo termina antes da assinatura, inclusive com textos longos.
                palavras = doc[0].get_text('words', clip=pymupdf.Rect(CAIXA_CORPO))
                self.assertTrue(all(p[3] <= CAIXA_CORPO[3] for p in palavras))
            imagem = pymupdf.Pixmap(str(convite.with_suffix('.png')))
            self.assertGreater(imagem.width, 500)
            self.assertEqual(len(list(saida.glob('*/*'))), 6)


if __name__ == "__main__":
    unittest.main()
