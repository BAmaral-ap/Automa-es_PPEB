"""Gera cartazes e convites de qualificação de doutorado a partir de uma planilha."""
from __future__ import annotations

import argparse
from datetime import date, datetime, time
import os
from pathlib import Path
import re
import shutil
import subprocess
import sys
import tempfile
import unicodedata
from zipfile import ZipFile

from docxtpl import DocxTemplate
from jinja2 import Environment, StrictUndefined
from lxml import etree
from openpyxl import load_workbook
from convite import FONTE_MAXIMA, conferir_pdf_convite, gerar_docx_convite
from layout import (CAIXAS, FONTE_MINIMA, aplicar_composicao, compor,
                    conferir_pdf, textos_caixas)

BASE = Path(__file__).resolve().parent
MEMBROS = (
    ("Presidente", "Universidade Presidente", "presidente", "Presidente"),
    ("Membro Interno", "Universidade Interno", "membro_interno", "Membro Interno"),
    ("Membro Externo 1", "Universidade Externo 1", "membro_externo_1", "Membro Externo"),
    ("Membro Externo 2", "Universidade Externo 2", "membro_externo_2", "Membro Externo"),
    ("Membro Externo 3", "Universidade Externo 3", "membro_externo_3", "Membro Externo"),
    ("Suplente 1", "Universidade Suplente 1", "suplente_1", "Suplente"),
    ("Suplente 2", "Universidade Suplente 2", "suplente_2", "Suplente"),
)
OBRIGATORIOS = ("Aluno", "Título", "Orientador", "Presidente", "Membro Interno",
                "Membro Externo 1", "Data", "Hora", "Sala")


def texto(valor):
    return "" if valor is None else str(valor).strip()


def formatar_data(valor):
    if isinstance(valor, (datetime, date)):
        return valor.strftime("%d/%m/%Y")
    for formato in ("%d/%m/%Y", "%Y-%m-%d"):
        try:
            return datetime.strptime(texto(valor), formato).strftime("%d/%m/%Y")
        except ValueError:
            pass
    raise ValueError("Data inválida; use DD/MM/AAAA ou uma data do Excel.")


def formatar_hora(valor):
    if isinstance(valor, datetime):
        valor = valor.time()
    if isinstance(valor, time):
        return valor.strftime("%Hh%M")
    if isinstance(valor, (int, float)) and not isinstance(valor, bool) and 0 <= valor < 1:
        minutos = round(valor * 24 * 60) % (24 * 60)
        return f"{minutos // 60:02d}h{minutos % 60:02d}"
    match = re.fullmatch(r"(\d{1,2})[:hH](\d{2})(?::00)?", texto(valor))
    if match and int(match[1]) < 24 and int(match[2]) < 60:
        return f"{int(match[1]):02d}h{int(match[2]):02d}"
    raise ValueError("Hora inválida; use HH:MM ou HHhMM.")


def nome_pasta(nome):
    nome = unicodedata.normalize("NFC", nome)
    nome = re.sub(r'[<>:"/\\|?*\x00-\x1f]', "_", nome).strip(" .")[:100].rstrip(" .")
    if not nome:
        raise ValueError("Nome do aluno não forma um nome de pasta válido.")
    if re.match(r"^(CON|PRN|AUX|NUL|COM[1-9]|LPT[1-9])(?:\.|$)", nome, re.I):
        nome = "Aluno - " + nome
    return nome


def contexto_linha(linha):
    faltam = [coluna for coluna in OBRIGATORIOS if not texto(linha.get(coluna))]
    if faltam:
        raise ValueError("Campos obrigatórios vazios: " + ", ".join(faltam))
    ctx = {"aluno": texto(linha["Aluno"]), "titulo": texto(linha["Título"]),
           "orientador": texto(linha["Orientador"]), "data": formatar_data(linha["Data"]),
           "hora": formatar_hora(linha["Hora"]), "sala": texto(linha["Sala"])}
    for coluna, universidade, chave, papel in MEMBROS:
        nome, uni = texto(linha.get(coluna)), texto(linha.get(universidade))
        if not nome and uni:
            raise ValueError(f"{universidade} preenchida sem {coluna}.")
        ctx[chave] = f"{nome}" + (f" ({uni})" if uni else "") + f" - {papel}" if nome else ""
    fonte = texto(linha.get("Fonte Título"))
    tamanho = float(fonte.replace(",", ".")) if fonte else 21
    if not 8 <= tamanho <= 36:
        raise ValueError("Fonte Título deve estar entre 8 e 36 pontos.")
    return ctx, tamanho


def ler_planilha(caminho):
    wb = load_workbook(caminho, data_only=False, read_only=True)
    try:
        ws = wb["Dados"] if "Dados" in wb.sheetnames else wb.active
        linhas = ws.iter_rows()
        headers = [texto(c.value) for c in next(linhas)]
        nomes = [h for h in headers if h]
        if len(nomes) != len(set(nomes)):
            raise ValueError("Há cabeçalhos duplicados na planilha.")
        faltam = set(OBRIGATORIOS) - set(headers)
        if faltam:
            raise ValueError("Colunas ausentes: " + ", ".join(sorted(faltam)))
        registros, pastas = [], set()
        for numero, cells in enumerate(linhas, 2):
            if all(c.value is None for c in cells):
                continue
            try:
                if any(c.data_type == "f" for c in cells):
                    raise ValueError("Use valores diretos, sem fórmulas.")
                linha = {h: c.value for h, c in zip(headers, cells) if h}
                ctx, fonte = contexto_linha(linha)
                pasta = nome_pasta(ctx["aluno"])
                if pasta.casefold() in pastas:
                    raise ValueError("Aluno duplicado ou nomes que geram a mesma pasta.")
                pastas.add(pasta.casefold())
                registros.append((numero, pasta, ctx, fonte))
            except (ValueError, TypeError) as exc:
                raise ValueError(f"Linha {numero}: {exc}") from exc
        if not registros:
            raise ValueError("A planilha não contém alunos para processar.")
        return registros
    finally:
        wb.close()


def gerar_docx(modelo, destino, contexto, fonte, limites=None):
    limites = limites or {}
    composicoes = {
        chave: compor(chave, textos, limites.get(chave, fonte if chave == "titulo" else CAIXAS[chave].fonte))
        for chave, textos in textos_caixas(contexto).items()
    }
    tpl = DocxTemplate(modelo)
    tpl.render(contexto, jinja_env=Environment(undefined=StrictUndefined), autoescape=True)
    tpl.save(destino)
    with ZipFile(destino) as arquivo:
        items = [(item, arquivo.read(item.filename)) for item in arquivo.infolist()]
    xml = etree.fromstring(dict((i.filename, data) for i, data in items)["word/document.xml"])
    aplicar_composicao(xml, composicoes)
    with ZipFile(destino, "w") as arquivo:
        for item, data in items:
            if item.filename == "word/document.xml":
                data = etree.tostring(xml, xml_declaration=True, encoding="UTF-8", standalone=True)
            arquivo.writestr(item, data)
    return composicoes


def localizar_soffice(explicito=None):
    candidatos = [explicito, os.environ.get("SOFFICE_PATH"), shutil.which("libreoffice"),
                  shutil.which("soffice"), "/Applications/LibreOffice.app/Contents/MacOS/soffice"]
    for var in ("PROGRAMFILES", "PROGRAMFILES(X86)"):
        if os.environ.get(var):
            candidatos.append(str(Path(os.environ[var]) / "LibreOffice/program/soffice.exe"))
    return next((str(Path(p).resolve()) for p in candidatos if p and Path(p).is_file()), None)


def escolher_conversor(nome, soffice):
    if nome != "word" and soffice:
        return "libreoffice"
    if nome != "libreoffice" and sys.platform in ("win32", "darwin"):
        return "word"
    raise ValueError("Conversão indisponível. Instale LibreOffice e use --soffice CAMINHO, "
                     "ou gere apenas DOCX com --formato docx. Word é suportado no Windows/macOS.")


def gerar_pdf(docx, conversor, soffice):
    destino = docx.with_suffix(".pdf")
    destino.unlink(missing_ok=True)
    if conversor == "libreoffice":
        with tempfile.TemporaryDirectory(prefix="ppeb-lo-") as profile:
            comando = [soffice, "-env:UserInstallation=" + Path(profile).as_uri(),
                       "--headless", "--convert-to", "pdf:writer_pdf_Export",
                       "--outdir", str(docx.parent), str(docx)]
            result = subprocess.run(comando, capture_output=True, text=True, timeout=120)
            if result.returncode or not destino.is_file():
                raise RuntimeError("LibreOffice não gerou o PDF: " + (result.stderr or result.stdout).strip())
    else:
        from docx2pdf import convert
        convert(str(docx), str(destino))
    if not destino.is_file() or destino.stat().st_size == 0:
        raise RuntimeError("O conversor não produziu um PDF válido.")
    return destino


def gerar_arquivos(modelo, docx, contexto, fonte, conversor, soffice, dpi):
    limites = {}
    # A composição por métricas também funciona no modo somente DOCX. Na
    # exportação, o próprio Word/LibreOffice confirma o que realmente ficou visível.
    for _ in range(57):
        composicoes = gerar_docx(modelo, docx, contexto, fonte, limites)
        if conversor is None:
            return
        pdf = gerar_pdf(docx, conversor, soffice)
        problemas = conferir_pdf(pdf, contexto)
        if not problemas:
            gerar_png(pdf, dpi)
            return
        for chave in problemas:
            if composicoes[chave].fonte <= FONTE_MINIMA:
                raise ValueError(f"O texto de '{chave}' não ficou integralmente visível "
                                 "na caixa fixa. Reduza o texto ou confira a fonte instalada.")
            limites[chave] = composicoes[chave].fonte - 0.5
    raise ValueError("Não foi possível ajustar o texto às caixas fixas.")


def gerar_png(pdf, dpi):
    import pymupdf
    with pymupdf.open(pdf) as doc:
        if len(doc) != 1:
            raise ValueError(f"Documento ficou com {len(doc)} páginas. Reduza o texto ou "
                             "ajuste o modelo antes de divulgar.")
        doc[0].get_pixmap(dpi=dpi, alpha=False).save(pdf.with_suffix(".png"))


def gerar_convite(modelo, docx, contexto, conversor, soffice, dpi):
    limite = FONTE_MAXIMA
    while limite >= FONTE_MINIMA:
        fonte, esperado = gerar_docx_convite(modelo, docx, contexto, limite)
        if conversor is None:
            return
        pdf = gerar_pdf(docx, conversor, soffice)
        if conferir_pdf_convite(pdf, esperado):
            gerar_png(pdf, dpi)
            return
        limite = fonte - 0.5
    raise ValueError("O texto do convite não ficou integralmente visível. "
                     "Reduza o texto ou confira a fonte instalada.")


def executar(args):
    registros = ler_planilha(args.planilha)
    if not args.modelo.is_file():
        raise ValueError(f"Modelo não encontrado: {args.modelo}")
    if not args.modelo_convite.is_file():
        raise ValueError(f"Modelo de convite não encontrado: {args.modelo_convite}")
    soffice = localizar_soffice(args.soffice)
    conversor = escolher_conversor(args.conversor, soffice) if args.formato == "todos" else None
    args.saida.mkdir(parents=True, exist_ok=True)
    for _, pasta, _, _ in registros:
        destino = args.saida / pasta
        if destino.exists() and (not args.sobrescrever or not destino.is_dir() or destino.is_symlink()):
            raise ValueError(f"Destino já existe: {destino}. Use outra saída ou --sobrescrever.")
    falhas = 0
    for numero, pasta, ctx, fonte in registros:
        try:
            with tempfile.TemporaryDirectory(prefix=".cartaz-", dir=args.saida) as staging:
                docx = Path(staging) / f"Cartaz - {pasta}.docx"
                gerar_arquivos(args.modelo, docx, ctx, fonte, conversor, soffice, args.dpi)
                convite = Path(staging) / f"Convite - {pasta}.docx"
                gerar_convite(args.modelo_convite, convite, ctx, conversor, soffice, args.dpi)
                destino = args.saida / pasta
                destino.mkdir(exist_ok=True)
                for arquivo in Path(staging).iterdir():
                    if args.sobrescrever:
                        os.replace(arquivo, destino / arquivo.name)
                    else:
                        shutil.move(str(arquivo), destino / arquivo.name)
            print(f"[OK] Linha {numero}: {destino}")
        except Exception as exc:
            falhas += 1
            print(f"[ERRO] Linha {numero}: {exc}", file=sys.stderr)
    print(f"Concluído: {len(registros) - falhas} conjunto(s) de cartaz e convite gerado(s); {falhas} falha(s).")
    return 1 if falhas else 0


def main(argv=None):
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--planilha", type=Path, default=BASE / "dados_qualificacao_doutorado.xlsx")
    parser.add_argument("--modelo", type=Path, default=BASE / "modelo_qualificacao_doutorado.docx")
    parser.add_argument("--modelo-convite", type=Path, default=BASE / "modelo_convite_doutorado.docx")
    parser.add_argument("--saida", type=Path, default=BASE / "cartazes_gerados")
    parser.add_argument("--formato", choices=("todos", "docx"), default="todos",
                        help="todos: cartaz e convite em DOCX, PDF e PNG; docx: os dois editáveis")
    parser.add_argument("--conversor", choices=("auto", "libreoffice", "word"), default="auto")
    parser.add_argument("--soffice", help="Caminho do executável LibreOffice")
    parser.add_argument("--dpi", type=int, default=300)
    parser.add_argument("--sobrescrever", action="store_true")
    args = parser.parse_args(argv)
    if not 72 <= args.dpi <= 600:
        parser.error("--dpi deve estar entre 72 e 600")
    args.saida = args.saida.resolve()
    try:
        return executar(args)
    except Exception as exc:
        print(f"[ERRO] {exc}", file=sys.stderr)
        return 1


if __name__ == "__main__":
    raise SystemExit(main())
